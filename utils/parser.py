"""
Module 1 (of the scoring pipeline): Document Parsing
Extracts raw text from uploaded resume files (PDF / DOCX).
Corresponds to Objective 2 and the "Text Extraction" stage of the methodology.
"""
import os
import pdfplumber
import docx


def extract_text_from_pdf(filepath):
    """Extract text from a PDF file using pdfplumber."""
    text_parts = []
    try:
        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
    except Exception as e:
        raise ValueError(f"Could not read PDF '{os.path.basename(filepath)}': {e}")
    return "\n".join(text_parts)


def extract_text_from_docx(filepath):
    """Extract text from a DOCX file using python-docx."""
    try:
        document = docx.Document(filepath)
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        # Also pull text out of any tables (common in resume templates)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)
        return "\n".join(paragraphs)
    except Exception as e:
        raise ValueError(f"Could not read DOCX '{os.path.basename(filepath)}': {e}")


def extract_text(filepath):
    """Dispatch to the right extractor based on file extension."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(filepath)
    elif ext == ".docx":
        return extract_text_from_docx(filepath)
    else:
        raise ValueError(f"Unsupported file type: {ext}. Only .pdf and .docx are accepted.")
